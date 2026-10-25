#!/usr/bin/env python
"""Extract `Platinum Marketing Plan.xlsx` into clean seed data for the audit system.

Outputs (default `data/`):
  channels.seed.json   contracts[] + channels[] for the current roster sheet
  channels.csv         one channel per row (for the marketing team)
  monthly_summary.csv  one row per monthly plan sheet (spend, counts, public/private)
  channel_history.csv  (month x primary channel) rows for roster-era sheets
  reg_by_month.csv     registrations per channel per month (from the Reg column)
  reg_list_monthly.csv registrations per source per month (from the Reg List sheet)
  channel_events.csv   stop / refund / follow-up / substitute events from side sheets
  data_quality.md      everything in the sheet that is ambiguous or inconsistent

Run with the scoop Python that has openpyxl:
  C:/Users/ADMIN/scoop/apps/python313/current/python -I tools/extract_channels.py
"""
from __future__ import annotations

import argparse
import collections
import csv
import datetime as dt
import json
import pathlib
import re
import sys

# --------------------------------------------------------------------------- #
# small parsers (unit-tested in tools/tests/test_extract.py)
# --------------------------------------------------------------------------- #

MONTHS = {
    "jan": 1, "january": 1, "feb": 2, "fab": 2, "february": 2, "mar": 3, "march": 3,
    "apr": 4, "april": 4, "may": 5, "jun": 6, "june": 6, "jul": 7, "july": 7,
    "aug": 8, "august": 8, "sep": 9, "sept": 9, "september": 9, "septemher": 9,
    "oct": 10, "october": 10, "nov": 11, "november": 11, "dec": 12, "december": 12,
}
WORDNUM = {"one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6, "seven": 7}
EXCLUSIVITY = {"us": 1, "two": 2, "three": 3, "thtee": 3, "four": 4, "five": 5,
               "2": 2, "3": 3, "4": 4, "5": 5}
SECTION_EXCLUSIVITY = {"only_us": 1, "only_two": 2, "only_three": 3, "only_four": 4, "only_five": 5}
ACTIVE_STATUSES = {"on process", "posted", "approved", "scheduled"}

REG_RE = re.compile(r"([A-Za-z]{3,9})\s*[Rr][eu]g\s*-?\s*(\d+)")
STREAK_RE = re.compile(r"(\d+)\s*(?:st|nd|rd|th)\b", re.I)
DATE_RE = re.compile(r"(\d{1,2})\.(\d{1,2})\.(\d{2,4})")
URL_RE = re.compile(r"https?://[^\s]+")
NUM_RE = re.compile(r"\d+(?:\.\d+)?\s*[kKmM]?(?![A-Za-z])")


def _s(v) -> str:
    return "" if v is None else str(v)


def parse_reg(text) -> list[dict]:
    """'Sep Reg -2 / Aug Reg -2' -> [{'month': 'Sep', 'reg': 2}, ...]"""
    if not text:
        return []
    return [{"month": m.group(1), "reg": int(m.group(2))} for m in REG_RE.finditer(str(text))]


def parse_streak(text) -> int | None:
    """'33th time continue' -> 33"""
    if not text:
        return None
    m = STREAK_RE.search(str(text))
    return int(m.group(1)) if m else None


def _iso(d: str, m: str, y: str):
    year = int(y)
    if year < 100:
        year += 2000
    try:
        return dt.date(year, int(m), int(d)).isoformat(), None
    except ValueError:
        return None, f"invalid_date:{d}.{m}.{y}"


def parse_period(text) -> dict:
    """'1.10.26 - 31.10.26' -> {'from': '2026-10-01', 'to': '2026-10-31', 'raw': ..., 'flags': []}"""
    raw = _s(text)
    out = {"from": None, "to": None, "raw": raw.strip(), "flags": []}
    dates = DATE_RE.findall(raw)
    if not dates:
        out["flags"].append("unparsed")
        return out
    out["from"], err = _iso(*dates[0])
    if err:
        out["flags"].append(err)
    if len(dates) >= 2:
        out["to"], err = _iso(*dates[1])
        if err:
            out["flags"].append(err)
    else:
        out["flags"].append("open_ended")
    return out


def parse_remark(text) -> dict:
    """Contract terms live in the Remark column: '1 day 1 post / pin daily / only us'."""
    res = {"quota_type": "unknown", "quota_posts_per_day": None, "pin_rule": "none", "exclusivity": None}
    if not text:
        return res
    s = re.sub(r"\s+", " ", str(text)).strip().lower()
    m = re.search(r"(\d+)\s*day\s*(\d+)\s*post", s)
    if m and int(m.group(1)) == 1:
        res["quota_type"] = "daily_post"
        res["quota_posts_per_day"] = int(m.group(2))
    else:
        res["quota_type"] = "custom"
    if re.search(r"pin\s*daily|daily\s*pin", s):
        res["pin_rule"] = "daily"
    elif (m := re.search(r"1\s*month\s*(\d+)\s*pin", s)):
        res["pin_rule"] = f"per_month:{m.group(1)}"
    elif (m := re.search(r"(\w+)\s*days?\s*a\s*pin", s)):
        n = WORDNUM.get(m.group(1), m.group(1))
        res["pin_rule"] = f"every_days:{n}"
    m = re.search(r"only\s*(us|two|three|thtee|four|five|[2-5])\b", s)
    if m:
        res["exclusivity"] = EXCLUSIVITY[m.group(1)]
    return res


def _clean_url(u: str) -> str:
    u = u.split("(")[0]
    u = u.strip().lstrip("\u200e\u200f\ufeff")
    return u.rstrip(")/.,;:'\"")


def classify_url(url: str) -> dict:
    info = {"kind": "external", "url": url, "username": None, "invite_hash": None, "sample_post_id": None}
    m = re.match(r"https?://(?:www\.)?t\.me/(.+)$", url, re.I)
    if not m:
        return info
    path = m.group(1)
    if path.startswith("+"):
        info.update(kind="private", invite_hash=path[1:].split("/")[0].split("?")[0])
        info["url"] = f"https://t.me/+{info['invite_hash']}"
        return info
    low = path.lower()
    if low.startswith("joinchat/"):
        info.update(kind="private", invite_hash=path.split("/", 1)[1].split("?")[0])
        return info
    if low.startswith("c/"):
        info.update(kind="private")
        return info
    if low.startswith("s/"):
        path = path[2:]
    parts = path.split("?")[0].split("/")
    username = parts[0]
    if not re.fullmatch(r"[A-Za-z0-9_]{3,}", username):
        info["kind"] = "other"
        return info
    post_id = int(parts[1]) if len(parts) > 1 and parts[1].isdigit() else None
    info.update(kind="public", url=f"https://t.me/{username}", username=username.lower(), sample_post_id=post_id)
    return info


def extract_links(cell_text, hyperlink_target=None) -> list[dict]:
    """All URLs in a cell (text + hyperlink target), cleaned, classified, de-duplicated."""
    cands = URL_RE.findall(str(cell_text)) if cell_text else []
    if hyperlink_target:
        cands.append(str(hyperlink_target))
    out, seen = [], set()
    for raw in cands:
        u = _clean_url(raw)
        if not u or "://" not in u:
            continue
        info = classify_url(u)
        key = info["username"] or info["invite_hash"] or info["url"].lower()
        if key in seen:
            continue
        seen.add(key)
        out.append(info)
    return out


def parse_k(s) -> int | None:
    """'5.7K' -> 5700, '829' -> 829"""
    t = str(s).strip().lower().replace(",", "")
    m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*([km])?", t)
    if not m:
        return None
    mult = {"k": 1000, "m": 1_000_000}.get(m.group(2), 1)
    return int(round(float(m.group(1)) * mult))


def parse_sub_view(text) -> dict:
    """'sub -23K/19K / 3 day ago view -290/210 / 7 day ago view -394/265' -> lists per metric."""
    res = {"subs": [], "views_3d": [], "views_7d": []}
    if not text:
        return res
    s = str(text).replace("\n", " ")
    anchors = []
    for key, pat in (("subs", r"\bsubs?\b"), ("views_3d", r"\b3\s*days?\b"), ("views_7d", r"\b7\s*days?\b")):
        for m in re.finditer(pat, s, re.I):
            anchors.append((m.start(), m.end(), key))
    anchors.sort()
    for i, (_, end, key) in enumerate(anchors):
        seg_end = anchors[i + 1][0] if i + 1 < len(anchors) else len(s)
        seg = re.sub(r"ago|views?|follower\w*|telegram|channel", " ", s[end:seg_end], flags=re.I)
        for tok in NUM_RE.findall(seg):
            v = parse_k(tok)
            if v is not None:
                res[key].append(v)
    return res


def section_from_header(text) -> str:
    s = re.sub(r"\s+", " ", _s(text)).strip().lower()
    m = re.match(r"only\s*(us|two|three|four|five)", s)
    return f"only_{m.group(1)}" if m else "influencer"


def split_names(text) -> list[str]:
    """Channel names in one cell: one per line. A ' / ' with whitespace around it also separates names,
    but a '/' inside parentheses or inside a word ('MM/Sub') is part of the name."""
    if not text:
        return []
    out = []
    for line in str(text).split("\n"):
        depth, buf, parts = 0, [], []
        for i, ch in enumerate(line):
            if ch in "([":
                depth += 1
            elif ch in ")]":
                depth = max(0, depth - 1)
            if ch == "/" and depth == 0:
                prev_ws = i == 0 or line[i - 1].isspace()
                next_ws = i == len(line) - 1 or line[i + 1].isspace()
                if prev_ws or next_ws:
                    parts.append("".join(buf))
                    buf = []
                    continue
            buf.append(ch)
        parts.append("".join(buf))
        out.extend(p.strip() for p in parts if p.strip())
    return out


def parse_admins(text) -> list[str]:
    if not text:
        return []
    return re.findall(r"@[A-Za-z0-9_]+", str(text))


def sheet_month_from_title(title: str) -> str | None:
    """'October 2026 MKT Pla' -> '2026-10'; first-era sheets without a year map to 2022-11 .. 2023-12."""
    t = _s(title).strip()
    if "mkt pla" not in t.lower():
        return None
    m = re.match(r"([A-Za-z]+)\s*z?\s*(\d{4}|\d{2})?\s*(\d)?\s*MKT", t, re.I)
    if not m:
        return None
    mon = MONTHS.get(m.group(1).lower())
    if not mon:
        return None
    if m.group(2):
        year = int(m.group(2))
        year = year + 2000 if year < 100 else year
    elif m.group(3):          # 'Nov 2 MKT Plan', 'Dec2 MKT Plan' = second Nov/Dec = 2023
        year = 2023
    else:                     # 'Nov MKT Plan' (2022) .. 'Oct MKT Plan' (2023)
        year = 2022 if mon >= 11 else 2023
    return f"{year:04d}-{mon:02d}"


HEADER_ALIASES = [
    ("reg", ("reg",)), ("admin", ("admin account",)), ("streak", ("continue time",)),
    ("period", ("plan to post date",)), ("name", ("cele name",)), ("link", ("page / acc",)),
    ("sub_view", ("post / ads active",)), ("status", ("situration",)), ("payment_note", ("payment note",)),
    ("continue_duration", ("continue duration",)), ("amount", ("total amount",)), ("payment1", ("payment 1",)),
    ("payment2", ("payment 2",)), ("payment_date", ("payment date",)), ("remark", ("remark",)),
    ("title", ("title",)), ("note", ("note",)), ("posted_date", ("posted date",)),
    ("post_link", ("post link",)), ("old_name", ("old name",)), ("old_link", ("old link",)), ("days", ("days",)),
]


def _norm(h) -> str:
    return re.sub(r"\s+", " ", _s(h)).strip().lower()


def map_headers(headers) -> dict:
    out = {}
    for i, h in enumerate(headers):
        n = _norm(h)
        if not n:
            continue
        for key, pats in HEADER_ALIASES:
            if key in out:
                continue
            if any(n.startswith(p) for p in pats):
                out[key] = i
                break
    return out


def find_header_row(ws, needle="situration", scan=12):
    for r in range(1, min(ws.max_row, scan) + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if any(_norm(v) == needle for v in vals):
            return r, vals
    raise ValueError(f"no header row with {needle!r} in sheet {ws.title!r}")


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    m = re.search(r"\d[\d,]*(?:\.\d+)?", str(v))
    if not m:
        return None
    try:
        f = float(m.group(0).replace(",", ""))
    except ValueError:
        return None
    return int(f) if f.is_integer() else f


def _blank(v) -> bool:
    return v is None or str(v).strip() == ""


def _text(v):
    if v is None:
        return None
    if isinstance(v, (dt.date, dt.datetime)):
        return v.isoformat()
    s = str(v).strip()
    return s or None


# --------------------------------------------------------------------------- #
# sheet-level extraction
# --------------------------------------------------------------------------- #

def parse_month_sheet(ws, month: str | None = None) -> list[dict]:
    month = month or sheet_month_from_title(ws.title) or "unknown"
    hdr_row, headers = find_header_row(ws)
    cols = map_headers(headers)
    if "name" not in cols or "link" not in cols:
        raise ValueError(f"sheet {ws.title!r} lacks name/link columns: {cols}")

    def cell(r, key):
        return ws.cell(r, cols[key] + 1) if key in cols else None

    def val(r, key):
        c = cell(r, key)
        return c.value if c is not None else None

    contracts, section = [], "unknown"
    for r in range(hdr_row + 1, ws.max_row + 1):
        name, link = val(r, "name"), val(r, "link")
        if _blank(name) and _blank(link):
            txt = " ".join(_s(ws.cell(r, c).value) for c in range(1, 8) if not _blank(ws.cell(r, c).value))
            if txt.strip():
                sec = section_from_header(txt)
                if sec != "influencer" or re.search(r"model|idol|page|creator|tiktok", txt, re.I):
                    section = sec
            continue
        link_cell = cell(r, "link")
        hl = link_cell.hyperlink.target if link_cell is not None and link_cell.hyperlink else None
        old_cell = cell(r, "old_link")
        old_hl = old_cell.hyperlink.target if old_cell is not None and old_cell.hyperlink else None
        old_links = extract_links(val(r, "old_link"), old_hl)
        old_keys = {_link_key(l) for l in old_links}
        links = extract_links(link, None)
        text_keys = {_link_key(l) for l in links}
        for l in extract_links(None, hl):
            k = _link_key(l)
            if k in text_keys or k in old_keys:      # a hyperlink still pointing at the Old Link stays old
                continue
            links.append(l)
            text_keys.add(k)
        old_links = [l for l in old_links if _link_key(l) not in text_keys]
        remark = parse_remark(val(r, "remark"))
        period = parse_period(val(r, "period"))
        amount_cell = cell(r, "amount")
        amount = _num(amount_cell.value) if amount_cell is not None else None
        names = split_names(name)
        tme_links = [l for l in links if l["kind"] in ("public", "private")]
        flags = list(period["flags"])
        if amount is None:
            flags.append("no_amount")
        if not links:
            flags.append("no_links")
        if len(names) > 1 and len(tme_links) > 1 and len(names) != len(tme_links):
            flags.append("names_links_mismatch")
        if any(l["kind"] == "other" for l in links):
            flags.append("unclassified_link")
        if remark["quota_type"] == "unknown":
            flags.append("no_remark")
        status = _text(val(r, "status"))
        primary = next((l for l in links if l["kind"] in ("public", "private")), links[0] if links else None)
        contracts.append({
            "contract_id": f"{month}-r{r:03d}",
            "month": month,
            "row": r,
            "section": section,
            "names": names,
            "admin_accounts": parse_admins(val(r, "admin")),
            "streak_months": parse_streak(val(r, "streak")),
            "period_from": period["from"],
            "period_to": period["to"],
            "period_raw": period["raw"],
            "quota_type": remark["quota_type"],
            "quota_posts_per_day": remark["quota_posts_per_day"],
            "pin_rule": remark["pin_rule"],
            "exclusivity": remark["exclusivity"] if remark["exclusivity"] is not None else SECTION_EXCLUSIVITY.get(section),
            "amount_mmk": amount,
            "payment1_mmk": _num(val(r, "payment1")),
            "payment2_mmk": _num(val(r, "payment2")),
            "payment_date_raw": _text(val(r, "payment_date")),
            "status_sheet": status,
            "reg_history": parse_reg(val(r, "reg")),
            "sub_view_raw": _text(val(r, "sub_view")),
            "sub_view": parse_sub_view(val(r, "sub_view")),
            "remark_raw": _text(val(r, "remark")),
            "payment_channel_raw": _text(val(r, "note")),
            "posted_date_raw": _text(val(r, "posted_date")),
            "old_name": _text(val(r, "old_name")),
            "old_links": old_links,
            "links": links,
            "primary_channel_id": channel_id(primary) if primary else None,
            "price_history_note": (amount_cell.comment.text.strip() if amount_cell is not None and amount_cell.comment else None),
            "flags": flags,
        })
    return contracts


def _link_key(l: dict) -> str:
    return l["username"] or l["invite_hash"] or l["url"].lower()


def channel_id(link: dict) -> str:
    if link["kind"] == "public":
        return f"pub:{link['username']}"
    if link["kind"] == "private" and link["invite_hash"]:
        return f"prv:{link['invite_hash']}"
    if link["kind"] == "private":
        return f"prvc:{link['url'].lower()}"
    return f"ext:{link['url'].lower()}"


def flatten_channels(contracts: list[dict]) -> list[dict]:
    out = []
    for c in contracts:
        names, links = c["names"], c["links"]
        positional = len(names) == len(links) and len(links) > 0
        for i, l in enumerate(links):
            nm = names[i] if positional else (names[0] if names else None)
            out.append(_channel_row(c, l, nm, False))
        for l in c["old_links"]:
            out.append(_channel_row(c, l, c.get("old_name") or (names[0] if names else None), True))
    return out


def _channel_row(c: dict, l: dict, name, is_old: bool) -> dict:
    return {
        "channel_id": channel_id(l),
        "contract_id": c["contract_id"],
        "month": c["month"],
        "kind": l["kind"],
        "username": l["username"],
        "invite_hash": l["invite_hash"],
        "url": l["url"],
        "sample_post_id": l["sample_post_id"],
        "name": name,
        "is_old_link": is_old,
        "section": c["section"],
        "exclusivity": c["exclusivity"],
        "quota_type": c["quota_type"],
        "quota_posts_per_day": c["quota_posts_per_day"],
        "pin_rule": c["pin_rule"],
        "amount_mmk_contract": c["amount_mmk"],
        "admin_accounts": c["admin_accounts"],
    }


def sheet_totals(ws, contracts: list[dict] | None = None) -> dict:
    hdr_row, headers = find_header_row(ws)
    cols = map_headers(headers)
    contracts = contracts if contracts is not None else parse_month_sheet(ws)

    def header_num(key):
        if key not in cols:
            return None
        for r in range(1, hdr_row):
            v = _num(ws.cell(r, cols[key] + 1).value)
            if v is not None:
                return v
        return None

    active = [c for c in contracts if (c["status_sheet"] or "").strip().lower() in ACTIVE_STATUSES]
    channels = flatten_channels(contracts)
    cur = [ch for ch in channels if not ch["is_old_link"]]
    status_counts = collections.Counter((c["status_sheet"] or "").strip() for c in contracts)
    return {
        "amount_header": header_num("amount"),
        "payment1_header": header_num("payment1"),
        "payment2_header": header_num("payment2"),
        "amount_sum": sum(c["amount_mmk"] or 0 for c in contracts),
        "amount_active_sum": sum(c["amount_mmk"] or 0 for c in active),
        "contracts": len(contracts),
        "active_contracts": len(active),
        "public_channels": len({ch["channel_id"] for ch in cur if ch["kind"] == "public"}),
        "private_channels": len({ch["channel_id"] for ch in cur if ch["kind"] == "private"}),
        "external_links": len({ch["channel_id"] for ch in cur if ch["kind"] == "external"}),
        "renamed_rows": sum(1 for c in contracts if c["old_name"] or c["old_links"]),
        "status_counts": dict(status_counts),
    }


def parse_reg_list(ws) -> list[dict]:
    """Reg List sheet: rows = sources, columns = month dates (row 2). -> long rows."""
    months = {}
    for c in range(2, ws.max_column + 1):
        v = ws.cell(2, c).value
        if isinstance(v, (dt.date, dt.datetime)):
            months[c] = f"{v.year:04d}-{v.month:02d}"
    rows = []
    for r in range(3, ws.max_row + 1):
        src = _text(ws.cell(r, 1).value)
        if not src:
            continue
        for c, ym in months.items():
            v = _num(ws.cell(r, c).value)
            if v is not None:
                rows.append({"source": src, "month": ym, "reg": v, "is_total": src.strip().lower() == "total"})
    return rows


def parse_side_sheet(ws, event: str) -> list[dict]:
    """Stop / Refund / Follow-up / Substitute sheets: pick columns by loose header match."""
    hdr = None
    for r in range(1, min(ws.max_row, 8) + 1):
        vals = [_norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if any("link" in v for v in vals) and any("name" in v for v in vals):
            hdr = r
            break
    if hdr is None:
        return []
    heads = [_norm(ws.cell(hdr, c).value) for c in range(1, ws.max_column + 1)]

    def col(sub):
        for i, h in enumerate(heads):
            if sub in h:
                return i + 1
        return None

    c_name, c_link, c_price, c_remark, c_time, c_admin, c_cond, c_refund = (
        col("name"), col("link"), col("price"), col("remark"), col("time"), col("admin"), col("contion"), col("refund"))
    out = []
    for r in range(hdr + 1, ws.max_row + 1):
        link_v = ws.cell(r, c_link).value if c_link else None
        name_v = ws.cell(r, c_name).value if c_name else None
        if _blank(link_v) and _blank(name_v):
            continue
        lc = ws.cell(r, c_link) if c_link else None
        links = extract_links(link_v, lc.hyperlink.target if lc is not None and lc.hyperlink else None)
        out.append({
            "event": event,
            "sheet": ws.title.strip(),
            "row": r,
            "name": _text(name_v),
            "channel_ids": [channel_id(l) for l in links],
            "urls": [l["url"] for l in links],
            "price_mmk": _num(ws.cell(r, c_price).value) if c_price else None,
            "refund_mmk": _num(ws.cell(r, c_refund).value) if c_refund else None,
            "condition_raw": _text(ws.cell(r, c_cond).value) if c_cond else None,
            "time_raw": _text(ws.cell(r, c_time).value) if c_time else None,
            "admin_raw": _text(ws.cell(r, c_admin).value) if c_admin else None,
            "remark_raw": _text(ws.cell(r, c_remark).value) if c_remark else None,
        })
    return out


# --------------------------------------------------------------------------- #
# whole-workbook pipeline
# --------------------------------------------------------------------------- #

def resolve_reg_months(sheet_month: str, reg_history: list[dict]) -> list[dict]:
    """Reg entries name a month only ('Sep'); resolve to YYYY-MM relative to the sheet month."""
    y, m = (int(x) for x in sheet_month.split("-"))
    out = []
    for e in reg_history:
        mon = MONTHS.get(e["month"].lower())
        if not mon:
            continue
        year = y if mon < m else y - 1
        out.append({"month": f"{year:04d}-{mon:02d}", "reg": e["reg"]})
    return out


def build_all(xlsx: pathlib.Path, roster_sheet: str, history_from: str):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    month_sheets = [(sheet_month_from_title(ws.title), ws) for ws in wb.worksheets]
    month_sheets = sorted(((m, ws) for m, ws in month_sheets if m), key=lambda t: t[0])

    monthly, history, reg_rows, errors = [], [], [], []
    roster_contracts = None
    for m, ws in month_sheets:
        try:
            contracts = parse_month_sheet(ws, m)
            tot = sheet_totals(ws, contracts)
        except ValueError as e:
            errors.append(str(e))
            continue
        monthly.append({"month": m, "sheet": ws.title.strip(), "era": "roster" if m >= history_from else "prospecting", **tot})
        if ws.title == roster_sheet:
            roster_contracts = contracts
        if m >= history_from:
            for c in contracts:
                history.append({
                    "month": m, "primary_channel_id": c["primary_channel_id"], "contract_id": c["contract_id"],
                    "name": " / ".join(c["names"]), "section": c["section"], "exclusivity": c["exclusivity"],
                    "quota_type": c["quota_type"], "pin_rule": c["pin_rule"], "amount_mmk": c["amount_mmk"],
                    "streak_months": c["streak_months"], "status": c["status_sheet"],
                    "public_links": sum(1 for l in c["links"] if l["kind"] == "public"),
                    "private_links": sum(1 for l in c["links"] if l["kind"] == "private"),
                    "admin_accounts": " ".join(c["admin_accounts"]),
                })
                for e in resolve_reg_months(m, c["reg_history"]):
                    reg_rows.append({"sheet_month": m, "primary_channel_id": c["primary_channel_id"],
                                     "contract_id": c["contract_id"], "reg_month": e["month"], "reg": e["reg"]})
    if roster_contracts is None:
        raise SystemExit(f"roster sheet {roster_sheet!r} not found")

    reg_list = parse_reg_list(wb["Reg List"]) if "Reg List" in wb.sheetnames else []
    events = []
    for title, ev in (("Stop list", "stop"), ("Refund ch", "refund"), ("Follow Up ch", "follow_up"),
                      (" Substitute ch", "substitute"), ("  Stop Ch ", "stop_ch"), ("break & Refund ", "break_refund")):
        if title in wb.sheetnames:
            events.extend(parse_side_sheet(wb[title], ev))
    return wb, roster_contracts, monthly, history, reg_rows, reg_list, events, errors


def data_quality_report(roster_sheet, contracts, channels, monthly, wb) -> str:
    cur = [ch for ch in channels if not ch["is_old_link"]]
    roster_tot = next(m for m in monthly if m["sheet"] == roster_sheet.strip())
    L = [f"# Data quality — `{roster_sheet}` (generated {dt.datetime.now():%Y-%m-%d %H:%M})", ""]
    L += ["## สรุป", "",
          f"- สัญญา (แถวที่มีชื่อหรือลิงก์): **{len(contracts)}**",
          f"- ยอดรวมคอลัมน์ Total amount: **{roster_tot['amount_sum']:,}** เทียบกับเซลล์หัวตาราง (J1): **{roster_tot['amount_header']:,}** → "
          + ("ตรงกัน" if roster_tot['amount_sum'] == roster_tot['amount_header'] else "**ไม่ตรง**"),
          f"- ช่องสาธารณะไม่ซ้ำ: **{roster_tot['public_channels']}**, ลิงก์เชิญส่วนตัวไม่ซ้ำ: **{roster_tot['private_channels']}**, ลิงก์ภายนอก (FB/TikTok): **{roster_tot['external_links']}**",
          f"- แถวที่มี Old Name/Old Link: **{roster_tot['renamed_rows']}**", ""]

    def section(title, rows):
        L.append(f"## {title} ({len(rows)})")
        L.append("")
        if not rows:
            L.append("- ไม่พบ")
        L.extend(rows)
        L.append("")

    section("ช่วงสัญญาที่ parse ไม่ได้หรือวันที่ผิด",
            [f"- `{c['contract_id']}` {c['names'][:1]} → `{c['period_raw']}` → {', '.join(f for f in c['flags'] if f.startswith(('invalid', 'unparsed', 'open')))}"
             for c in contracts if any(f.startswith(("invalid", "unparsed", "open_ended")) for f in c["flags"])])
    section("สัญญาที่ไม่มีราคา (Total amount ว่าง)",
            [f"- `{c['contract_id']}` {' / '.join(c['names'])[:60]} — remark: {(c['remark_raw'] or '')[:80]}" for c in contracts if "no_amount" in c["flags"]])
    section("จำนวนชื่อช่องไม่เท่ากับจำนวนลิงก์ (จับคู่ชื่อ-ลิงก์ไม่ได้ ใช้ชื่อแรกแทน)",
            [f"- `{c['contract_id']}` names={len(c['names'])} tme_links={sum(1 for l in c['links'] if l['kind'] in ('public', 'private'))}" for c in contracts if "names_links_mismatch" in c["flags"]])
    section("Remark ที่ไม่ใช่รูปแบบ `N day N post` (quota_type=custom — ต้องกำหนดกติกาเอง)",
            [f"- `{c['contract_id']}` {' / '.join(c['names'])[:50]} — {(c['remark_raw'] or '')[:100]}" for c in contracts if c["quota_type"] == "custom"])
    dup = collections.defaultdict(list)
    for ch in cur:
        dup[ch["channel_id"]].append(ch["contract_id"])
    section("channel_id ซ้ำข้ามสัญญาในเดือนเดียว",
            [f"- `{k}` ใน {v}" for k, v in dup.items() if len(set(v)) > 1])
    adm = collections.defaultdict(set)
    for c in contracts:
        for a in c["admin_accounts"]:
            adm[a.lower()].add(c["contract_id"])
    section("Admin account เดียวดูแลหลายสัญญา (ข้อมูลประกอบ ไม่ใช่ข้อผิดพลาด)",
            [f"- `{a}` → {len(v)} สัญญา" for a, v in sorted(adm.items(), key=lambda kv: -len(kv[1])) if len(v) > 1])
    section("ลิงก์ที่จัดประเภทไม่ได้",
            [f"- `{c['contract_id']}` {[l['url'] for l in c['links'] if l['kind'] == 'other']}" for c in contracts if "unclassified_link" in c["flags"]])
    section("สัญญาที่ไม่มีลิงก์เลย", [f"- `{c['contract_id']}` {c['names']}" for c in contracts if "no_links" in c["flags"]])
    section("Payment Date ว่าง", [f"- `{c['contract_id']}`" for c in contracts if not c["payment_date_raw"]])
    section("ลิงก์ t.me/c/… (private permalink ต้องเป็นสมาชิกจึงอ่านได้)",
            [f"- `{ch['contract_id']}` {ch['url']}" for ch in cur if ch["channel_id"].startswith("prvc:")])
    L += ["## หมายเหตุอื่น", "",
          "- เซลล์ H1 ของทุกชีตตั้งแต่ Oct 2023 มีค่า `129` คงที่ ไม่ทราบที่มา (ไม่ใช่สูตร) — ไม่ได้ใช้ในระบบใหม่",
          "- สูตรในชีตเดือนมีแค่ `=SUM` ที่ J1/K1/L1; ค่า sub/view ทั้งหมดพิมพ์มือ จึงใช้เป็น baseline เท่านั้น",
          "- ชีต 2022–2024 เป็นลิสต์ prospect ปน Rejected; `monthly_summary.csv` แยก `amount_active_sum` (On Process/Posted/Approved/Scheduled) ออกจาก `amount_sum`",
          ""]
    return "\n".join(L)


def write_outputs(out: pathlib.Path, roster_sheet, contracts, channels, monthly, history, reg_rows, reg_list, events, dq):
    out.mkdir(parents=True, exist_ok=True)
    seed = {"generated_at": dt.datetime.now().isoformat(timespec="seconds"), "source_sheet": roster_sheet,
            "contracts": contracts, "channels": channels}
    (out / "channels.seed.json").write_text(json.dumps(seed, ensure_ascii=False, indent=1), encoding="utf-8")

    def write_csv(name, rows, fields):
        with (out / name).open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
            w.writeheader()
            for r in rows:
                w.writerow({k: (json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v) for k, v in r.items()})

    write_csv("channels.csv", channels, list(channels[0].keys()) if channels else ["channel_id"])
    write_csv("monthly_summary.csv", monthly, list(monthly[0].keys()))
    write_csv("channel_history.csv", history, list(history[0].keys()) if history else ["month"])
    write_csv("reg_by_month.csv", reg_rows, ["sheet_month", "primary_channel_id", "contract_id", "reg_month", "reg"])
    write_csv("reg_list_monthly.csv", reg_list, ["source", "month", "reg", "is_total"])
    write_csv("channel_events.csv", events, list(events[0].keys()) if events else ["event"])
    (out / "data_quality.md").write_text(dq, encoding="utf-8")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", default="Platinum Marketing Plan.xlsx")
    ap.add_argument("--out", default="data")
    ap.add_argument("--roster", default="October 2026 MKT Pla")
    ap.add_argument("--history-from", default="2025-01")
    a = ap.parse_args(argv)
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    wb, contracts, monthly, history, reg_rows, reg_list, events, errors = build_all(
        pathlib.Path(a.xlsx), a.roster, a.history_from)
    channels = flatten_channels(contracts)
    dq = data_quality_report(a.roster, contracts, channels, monthly, wb)
    write_outputs(pathlib.Path(a.out), a.roster, contracts, channels, monthly, history, reg_rows, reg_list, events, dq)
    cur = [c for c in channels if not c["is_old_link"]]
    print(f"roster={a.roster!r} contracts={len(contracts)} amount_sum={sum(c['amount_mmk'] or 0 for c in contracts):,}")
    print(f"channels: public={len({c['channel_id'] for c in cur if c['kind']=='public'})} "
          f"private={len({c['channel_id'] for c in cur if c['kind']=='private'})} "
          f"external={len({c['channel_id'] for c in cur if c['kind']=='external'})} old_links={len(channels)-len(cur)}")
    print(f"monthly_summary rows={len(monthly)} history rows={len(history)} reg rows={len(reg_rows)} "
          f"reg_list rows={len(reg_list)} events={len(events)}")
    for e in errors:
        print("SKIPPED:", e)


if __name__ == "__main__":
    main()
